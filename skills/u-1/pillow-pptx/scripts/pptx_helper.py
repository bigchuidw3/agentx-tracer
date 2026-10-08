"""
PPTX Helper - 富文本PPT生成增强工具

提供丰富的PPT页面生成功能，支持背景图、图标、插图等元素。

依赖: python-pptx, Pillow
安装: pip install python-pptx Pillow
"""

import os
from typing import List, Dict, Optional, Any, Tuple, Union

try:
    from pptx import Presentation
    from pptx.util import Inches, Pt, Cm, Emu
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
except ImportError:
    raise ImportError("python-pptx is required. Install with: pip install python-pptx")

try:
    from PIL import Image
    PIL_AVAILABLE = True
except ImportError:
    PIL_AVAILABLE = False


# ==================== 配色方案 ====================

COLOR_SCHEMES = {
    '商务蓝': {
        'primary': RGBColor(27, 58, 92),
        'secondary': RGBColor(232, 240, 254),
        'accent': RGBColor(255, 107, 53),
        'background': RGBColor(245, 247, 250),
        'text': RGBColor(44, 62, 80),
        'light_text': RGBColor(255, 255, 255),
        'bg_dark': RGBColor(27, 58, 92),
    },
    '科技紫': {
        'primary': RGBColor(74, 29, 142),
        'secondary': RGBColor(240, 230, 255),
        'accent': RGBColor(0, 212, 255),
        'background': RGBColor(13, 13, 26),
        'text': RGBColor(224, 224, 224),
        'light_text': RGBColor(255, 255, 255),
        'bg_dark': RGBColor(13, 13, 26),
    },
    '清新绿': {
        'primary': RGBColor(45, 125, 70),
        'secondary': RGBColor(232, 245, 233),
        'accent': RGBColor(255, 183, 77),
        'background': RGBColor(241, 248, 233),
        'text': RGBColor(55, 71, 79),
        'light_text': RGBColor(255, 255, 255),
        'bg_dark': RGBColor(45, 125, 70),
    },
    '简约灰': {
        'primary': RGBColor(55, 71, 79),
        'secondary': RGBColor(236, 239, 241),
        'accent': RGBColor(255, 112, 67),
        'background': RGBColor(250, 250, 250),
        'text': RGBColor(55, 71, 79),
        'light_text': RGBColor(255, 255, 255),
        'bg_dark': RGBColor(55, 71, 79),
    },
    '创意橙': {
        'primary': RGBColor(230, 81, 0),
        'secondary': RGBColor(255, 243, 224),
        'accent': RGBColor(124, 77, 255),
        'background': RGBColor(255, 248, 225),
        'text': RGBColor(78, 52, 46),
        'light_text': RGBColor(255, 255, 255),
        'bg_dark': RGBColor(230, 81, 0),
    },
    '暗色系': {
        'primary': RGBColor(187, 134, 252),
        'secondary': RGBColor(30, 30, 46),
        'accent': RGBColor(3, 218, 198),
        'background': RGBColor(18, 18, 18),
        'text': RGBColor(224, 224, 224),
        'light_text': RGBColor(255, 255, 255),
        'bg_dark': RGBColor(18, 18, 18),
    }
}

# 默认字体
DEFAULT_FONT = '微软雅黑'
DEFAULT_FONT_FALLBACK = '思源黑体'


# ==================== 演示文稿创建 ====================

def create_presentation(
    title: str = "演示文稿",
    author: str = "Rich PPTX",
    subject: str = "",
    aspect_ratio: str = "16:9"
) -> Presentation:
    """
    创建演示文稿。

    Args:
        title: 标题
        author: 作者
        subject: 主题
        aspect_ratio: 宽高比 ("16:9" 或 "4:3")

    Returns:
        Presentation对象
    """
    prs = Presentation()

    if aspect_ratio == "16:9":
        prs.slide_width = Inches(10)
        prs.slide_height = Inches(5.625)
    elif aspect_ratio == "4:3":
        prs.slide_width = Inches(10)
        prs.slide_height = Inches(7.5)
    else:
        raise ValueError("aspect_ratio must be '16:9' or '4:3'")

    prs.core_properties.title = title
    prs.core_properties.author = author
    prs.core_properties.subject = subject

    return prs


# ==================== 辅助函数 ====================

def set_text_style(text_frame: Any, size: int, bold: bool = False,
                   color: RGBColor = None, alignment: PP_ALIGN = PP_ALIGN.LEFT) -> None:
    """设置文本样式"""
    for paragraph in text_frame.paragraphs:
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.alignment = alignment
        if color:
            paragraph.font.color.rgb = color


def add_background_image(slide: Any, image_path: str, opacity: float = 1.0) -> None:
    """为幻灯片添加背景图片"""
    if not os.path.exists(image_path):
        print(f"Warning: Background image not found: {image_path}")
        return

    try:
        slide.shapes.add_picture(
            image_path,
            Inches(0), Inches(0),
            width=Inches(10),
            height=Inches(5.625)
        )
    except Exception as e:
        print(f"Warning: Failed to add background image: {e}")


def set_slide_background_color(slide: Any, color: RGBColor) -> None:
    """设置幻灯片背景颜色"""
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_textbox(slide: Any, left: float, top: float, width: float, height: float,
                text: str, font_size: int = 18, bold: bool = False,
                color: RGBColor = None, alignment: PP_ALIGN = PP_ALIGN.LEFT) -> Any:
    """添加文本框"""
    shape = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = shape.text_frame
    tf.word_wrap = True
    tf.text = text

    for paragraph in tf.paragraphs:
        paragraph.font.size = Pt(font_size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = color if color else RGBColor(0, 0, 0)
        paragraph.alignment = alignment

    return shape


# ==================== 幻灯片类型 ====================

def add_rich_title_slide(
    prs: Presentation,
    title: str,
    subtitle: str = "",
    bg_image: str = None,
    decoration_image: str = None,
    style: str = '科技紫'
) -> Any:
    """
    添加富文本封面页（带背景和装饰）。

    Args:
        prs: Presentation对象
        title: 主标题
        subtitle: 副标题
        bg_image: 背景图片路径
        decoration_image: 装饰图片路径
        style: 风格名称

    Returns:
        幻灯片对象
    """
    colors = COLOR_SCHEMES.get(style, COLOR_SCHEMES['科技紫'])

    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout

    # 设置背景
    if bg_image and os.path.exists(bg_image):
        add_background_image(slide, bg_image)
    else:
        set_slide_background_color(slide, colors['background'])

    # 添加装饰图片
    if decoration_image and os.path.exists(decoration_image):
        try:
            pic = slide.shapes.add_picture(
                decoration_image,
                Inches(6.5), Inches(2),
                width=Inches(3.5)
            )
        except Exception as e:
            print(f"Warning: Failed to add decoration image: {e}")

    # 添加主标题
    title_shape = slide.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(8), Inches(1.2))
    title_tf = title_shape.text_frame
    title_tf.text = title
    for p in title_tf.paragraphs:
        p.font.size = Pt(44)
        p.font.bold = True
        p.font.color.rgb = colors['light_text']
        p.font.name = DEFAULT_FONT

    # 添加副标题
    if subtitle:
        subtitle_shape = slide.shapes.add_textbox(Inches(0.8), Inches(3.2), Inches(8), Inches(0.8))
        subtitle_tf = subtitle_shape.text_frame
        subtitle_tf.text = subtitle
        for p in subtitle_tf.paragraphs:
            p.font.size = Pt(24)
            p.font.color.rgb = colors['text']
            p.font.name = DEFAULT_FONT

    return slide


def add_icon_content_slide(
    prs: Presentation,
    title: str,
    items: List[Dict[str, str]],
    bg_image: str = None,
    style: str = '科技紫'
) -> Any:
    """
    添加图标内容页（图标+文字列表）。

    Args:
        prs: Presentation对象
        title: 页面标题
        items: 内容项列表，每项包含:
               - icon: 图标图片路径
               - title: 要点标题
               - desc: 要点描述
        bg_image: 背景图片路径
        style: 风格名称

    Returns:
        幻灯片对象
    """
    colors = COLOR_SCHEMES.get(style, COLOR_SCHEMES['科技紫'])

    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout

    # 设置背景
    if bg_image and os.path.exists(bg_image):
        add_background_image(slide, bg_image)
    else:
        set_slide_background_color(slide, colors['background'])

    # 添加页面标题
    title_shape = slide.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(9), Inches(0.8))
    title_tf = title_shape.text_frame
    title_tf.text = title
    for p in title_tf.paragraphs:
        p.font.size = Pt(32)
        p.font.bold = True
        p.font.color.rgb = colors['primary']
        p.font.name = DEFAULT_FONT

    # 计算布局
    item_count = len(items)
    if item_count > 4:
        item_count = 4

    start_y = 1.5
    item_height = 0.9
    icon_size = 0.6

    for i, item in enumerate(items[:4]):
        y_pos = start_y + i * item_height

        # 添加图标
        icon_path = item.get('icon')
        if icon_path and os.path.exists(icon_path):
            try:
                slide.shapes.add_picture(
                    icon_path,
                    Inches(0.6), Inches(y_pos),
                    width=Inches(icon_size)
                )
            except Exception as e:
                print(f"Warning: Failed to add icon: {e}")

        # 添加标题
        item_title = item.get('title', '')
        title_text = slide.shapes.add_textbox(Inches(1.4), Inches(y_pos), Inches(3), Inches(0.4))
        title_tf = title_text.text_frame
        title_tf.text = item_title
        for p in title_tf.paragraphs:
            p.font.size = Pt(18)
            p.font.bold = True
            p.font.color.rgb = colors['primary']
            p.font.name = DEFAULT_FONT

        # 添加描述
        desc = item.get('desc', '')
        if desc:
            desc_text = slide.shapes.add_textbox(Inches(1.4), Inches(y_pos + 0.35), Inches(7), Inches(0.5))
            desc_tf = desc_text.text_frame
            desc_tf.text = desc
            for p in desc_tf.paragraphs:
                p.font.size = Pt(14)
                p.font.color.rgb = colors['text']
                p.font.name = DEFAULT_FONT

    return slide


def add_illustrated_slide(
    prs: Presentation,
    title: str,
    content: Union[str, List[str]],
    illustration_image: str = None,
    layout: str = 'left',
    bg_image: str = None,
    style: str = '科技紫'
) -> Any:
    """
    添加插图内容页（左文右图或右文左图）。

    Args:
        prs: Presentation对象
        title: 页面标题
        content: 内容文本（字符串或字符串列表）
        illustration_image: 插图路径
        layout: 布局 ('left'=左文右图, 'right'=左图右文)
        bg_image: 背景图片路径
        style: 风格名称

    Returns:
        幻灯片对象
    """
    colors = COLOR_SCHEMES.get(style, COLOR_SCHEMES['科技紫'])

    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout

    # 设置背景
    if bg_image and os.path.exists(bg_image):
        add_background_image(slide, bg_image)
    else:
        set_slide_background_color(slide, colors['background'])

    # 布局
    if layout == 'left':
        text_left, text_width = 0.5, 5
        illust_left = 6
    else:
        text_left, text_width = 5.5, 4
        illust_left = 0.5

    # 添加标题
    title_shape = slide.shapes.add_textbox(Inches(text_left), Inches(0.5), Inches(text_width), Inches(0.8))
    title_tf = title_shape.text_frame
    title_tf.text = title
    for p in title_tf.paragraphs:
        p.font.size = Pt(32)
        p.font.bold = True
        p.font.color.rgb = colors['primary']
        p.font.name = DEFAULT_FONT

    # 添加内容
    content_y = 1.5
    if isinstance(content, str):
        content_items = [content]
    else:
        content_items = content

    for item in content_items:
        content_shape = slide.shapes.add_textbox(
            Inches(text_left), Inches(content_y),
            Inches(text_width), Inches(0.5)
        )
        content_tf = content_shape.text_frame
        content_tf.text = f"• {item}"
        for p in content_tf.paragraphs:
            p.font.size = Pt(16)
            p.font.color.rgb = colors['text']
            p.font.name = DEFAULT_FONT
        content_y += 0.5

    # 添加插图
    if illustration_image and os.path.exists(illustration_image):
        try:
            slide.shapes.add_picture(
                illustration_image,
                Inches(illust_left), Inches(1.5),
                width=Inches(3.5)
            )
        except Exception as e:
            print(f"Warning: Failed to add illustration: {e}")

    return slide


def add_chart_slide(
    prs: Presentation,
    title: str,
    chart_type: str,
    categories: List[str],
    series_data: Union[List[float], Dict[str, List[float]]],
    bg_image: str = None,
    style: str = '科技紫'
) -> Any:
    """
    添加图表页。

    Args:
        prs: Presentation对象
        title: 页面标题
        chart_type: 图表类型 ('bar', 'column', 'line', 'pie')
        categories: 分类标签列表
        series_data: 数据（单系列列表或多系列字典）
        bg_image: 背景图片路径
        style: 风格名称

    Returns:
        幻灯片对象
    """
    colors = COLOR_SCHEMES.get(style, COLOR_SCHEMES['科技紫'])

    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout

    # 设置背景
    if bg_image and os.path.exists(bg_image):
        add_background_image(slide, bg_image)
    else:
        set_slide_background_color(slide, colors['background'])

    # 添加标题
    title_shape = slide.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(9), Inches(0.8))
    title_tf = title_shape.text_frame
    title_tf.text = title
    for p in title_tf.paragraphs:
        p.font.size = Pt(32)
        p.font.bold = True
        p.font.color.rgb = colors['primary']
        p.font.name = DEFAULT_FONT

    # 图表类型映射
    chart_types = {
        'bar': XL_CHART_TYPE.BAR_CLUSTERED,
        'column': XL_CHART_TYPE.COLUMN_CLUSTERED,
        'line': XL_CHART_TYPE.LINE,
        'pie': XL_CHART_TYPE.PIE,
    }

    xl_chart_type = chart_types.get(chart_type.lower(), XL_CHART_TYPE.COLUMN_CLUSTERED)

    # 准备数据
    chart_data = CategoryChartData()
    chart_data.categories = categories

    if isinstance(series_data, dict):
        for series_name, values in series_data.items():
            chart_data.add_series(series_name, values)
    else:
        chart_data.add_series('数据', series_data)

    # 添加图表
    chart = slide.shapes.add_chart(
        xl_chart_type,
        Inches(1), Inches(1.5),
        Inches(8), Inches(4),
        chart_data
    ).chart

    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout = False

    return slide


def add_quote_slide(
    prs: Presentation,
    quote: str,
    author: str = "",
    bg_image: str = None,
    style: str = '科技紫'
) -> Any:
    """
    添加引用页（大字引言+装饰）。

    Args:
        prs: Presentation对象
        quote: 引言内容
        author: 作者
        bg_image: 背景图片路径
        style: 风格名称

    Returns:
        幻灯片对象
    """
    colors = COLOR_SCHEMES.get(style, COLOR_SCHEMES['科技紫'])

    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout

    # 设置背景
    if bg_image and os.path.exists(bg_image):
        add_background_image(slide, bg_image)
    else:
        set_slide_background_color(slide, colors['primary'])

    # 添加引言
    quote_shape = slide.shapes.add_textbox(Inches(1), Inches(1.8), Inches(8), Inches(2))
    quote_tf = quote_shape.text_frame
    quote_tf.word_wrap = True
    quote_tf.text = f"「{quote}」"
    for p in quote_tf.paragraphs:
        p.font.size = Pt(32)
        p.font.color.rgb = colors['light_text']
        p.font.name = DEFAULT_FONT
        p.alignment = PP_ALIGN.CENTER

    # 添加作者
    if author:
        author_shape = slide.shapes.add_textbox(Inches(1), Inches(4), Inches(8), Inches(0.5))
        author_tf = author_shape.text_frame
        author_tf.text = f"— {author}"
        for p in author_tf.paragraphs:
            p.font.size = Pt(18)
            p.font.color.rgb = colors['secondary']
            p.font.name = DEFAULT_FONT
            p.alignment = PP_ALIGN.RIGHT

    return slide


def add_section_slide(
    prs: Presentation,
    section_title: str,
    section_number: str = "",
    bg_image: str = None,
    style: str = '科技紫'
) -> Any:
    """
    添加章节分隔页。

    Args:
        prs: Presentation对象
        section_title: 章节标题
        section_number: 章节编号
        bg_image: 背景图片路径
        style: 风格名称

    Returns:
        幻灯片对象
    """
    colors = COLOR_SCHEMES.get(style, COLOR_SCHEMES['科技紫'])

    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout

    # 设置背景
    if bg_image and os.path.exists(bg_image):
        add_background_image(slide, bg_image)
    else:
        set_slide_background_color(slide, colors['primary'])

    # 添加章节编号
    if section_number:
        num_shape = slide.shapes.add_textbox(Inches(0.5), Inches(2), Inches(9), Inches(1))
        num_tf = num_shape.text_frame
        num_tf.text = section_number
        for p in num_tf.paragraphs:
            p.font.size = Pt(72)
            p.font.bold = True
            p.font.color.rgb = colors['accent']
            p.font.name = DEFAULT_FONT
            p.alignment = PP_ALIGN.CENTER

    # 添加章节标题
    title_shape = slide.shapes.add_textbox(Inches(0.5), Inches(3.2), Inches(9), Inches(1))
    title_tf = title_shape.text_frame
    title_tf.text = section_title
    for p in title_tf.paragraphs:
        p.font.size = Pt(40)
        p.font.bold = True
        p.font.color.rgb = colors['light_text']
        p.font.name = DEFAULT_FONT
        p.alignment = PP_ALIGN.CENTER

    return slide


def add_end_slide(
    prs: Presentation,
    title: str = "谢谢",
    subtitle: str = "",
    bg_image: str = None,
    style: str = '科技紫'
) -> Any:
    """
    添加结尾页。

    Args:
        prs: Presentation对象
        title: 结束语
        subtitle: 副标题/联系方式
        bg_image: 背景图片路径
        style: 风格名称

    Returns:
        幻灯片对象
    """
    colors = COLOR_SCHEMES.get(style, COLOR_SCHEMES['科技紫'])

    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout

    # 设置背景
    if bg_image and os.path.exists(bg_image):
        add_background_image(slide, bg_image)
    else:
        set_slide_background_color(slide, colors['background'])

    # 添加主标题
    title_shape = slide.shapes.add_textbox(Inches(0.5), Inches(2), Inches(9), Inches(1.2))
    title_tf = title_shape.text_frame
    title_tf.text = title
    for p in title_tf.paragraphs:
        p.font.size = Pt(56)
        p.font.bold = True
        p.font.color.rgb = colors['primary']
        p.font.name = DEFAULT_FONT
        p.alignment = PP_ALIGN.CENTER

    # 添加副标题
    if subtitle:
        subtitle_shape = slide.shapes.add_textbox(Inches(0.5), Inches(3.5), Inches(9), Inches(0.8))
        subtitle_tf = subtitle_shape.text_frame
        subtitle_tf.text = subtitle
        for p in subtitle_tf.paragraphs:
            p.font.size = Pt(20)
            p.font.color.rgb = colors['text']
            p.font.name = DEFAULT_FONT
            p.alignment = PP_ALIGN.CENTER

    return slide


def add_bullet_slide(
    prs: Presentation,
    title: str,
    bullets: List[str],
    bg_image: str = None,
    style: str = '科技紫'
) -> Any:
    """
    添加要点列表页。

    Args:
        prs: Presentation对象
        title: 页面标题
        bullets: 要点列表
        bg_image: 背景图片路径
        style: 风格名称

    Returns:
        幻灯片对象
    """
    colors = COLOR_SCHEMES.get(style, COLOR_SCHEMES['科技紫'])

    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout

    # 设置背景
    if bg_image and os.path.exists(bg_image):
        add_background_image(slide, bg_image)
    else:
        set_slide_background_color(slide, colors['background'])

    # 添加标题
    title_shape = slide.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(9), Inches(0.8))
    title_tf = title_shape.text_frame
    title_tf.text = title
    for p in title_tf.paragraphs:
        p.font.size = Pt(32)
        p.font.bold = True
        p.font.color.rgb = colors['primary']
        p.font.name = DEFAULT_FONT

    # 添加要点
    start_y = 1.5
    for i, bullet in enumerate(bullets):
        bullet_shape = slide.shapes.add_textbox(
            Inches(0.8), Inches(start_y + i * 0.6),
            Inches(8.5), Inches(0.5)
        )
        bullet_tf = bullet_shape.text_frame
        bullet_tf.text = f"• {bullet}"
        for p in bullet_tf.paragraphs:
            p.font.size = Pt(18)
            p.font.color.rgb = colors['text']
            p.font.name = DEFAULT_FONT

    return slide


# ==================== 保存功能 ====================

def save_presentation(
    prs: Presentation,
    filename: str,
    verbose: bool = True
) -> str:
    """
    保存演示文稿。

    Args:
        prs: Presentation对象
        filename: 输出文件名（可以是绝对路径或相对路径）
        verbose: 是否打印详细信息

    Returns:
        保存文件的绝对路径
    """
    # 确保目录存在
    output_dir = os.path.dirname(filename)
    if output_dir and not os.path.exists(output_dir):
        os.makedirs(output_dir)

    # 确保.pptx扩展名
    if not filename.endswith('.pptx'):
        filename += '.pptx'

    # 保存
    prs.save(filename)

    # 获取绝对路径
    abs_path = os.path.abspath(filename)

    if verbose:
        file_size = os.path.getsize(abs_path) / (1024 * 1024)
        print(f"演示文稿已保存: {abs_path}")
        print(f"  页数: {len(prs.slides)}")
        print(f"  大小: {file_size:.2f} MB")

    return abs_path


# ==================== 工具函数 ====================

def get_color_scheme(style_name: str) -> Dict[str, RGBColor]:
    """获取指定风格的配色"""
    return COLOR_SCHEMES.get(style_name, COLOR_SCHEMES['科技紫'])


def list_color_schemes() -> List[str]:
    """列出所有可用的配色方案"""
    return list(COLOR_SCHEMES.keys())


if __name__ == "__main__":
    print("PPTX Helper Module")
    print("=" * 50)
    print("\nAvailable functions:")
    print("  - create_presentation(title, author, aspect_ratio)")
    print("  - add_rich_title_slide(prs, title, subtitle, bg_image, decoration_image)")
    print("  - add_icon_content_slide(prs, title, items, bg_image)")
    print("  - add_illustrated_slide(prs, title, content, illustration_image, layout)")
    print("  - add_chart_slide(prs, title, chart_type, categories, series_data)")
    print("  - add_quote_slide(prs, quote, author)")
    print("  - add_section_slide(prs, section_title, section_number)")
    print("  - add_end_slide(prs, title, subtitle)")
    print("  - add_bullet_slide(prs, title, bullets)")
    print("  - save_presentation(prs, filename)")
    print("\nColor schemes:", list_color_schemes())
